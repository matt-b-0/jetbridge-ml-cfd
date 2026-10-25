from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).parent/'report_output'
OUT.mkdir(exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.2677); sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.78)
sec.header_distance=sec.footer_distance=Inches(.28)
for name,size in [('Normal',10.5),('Title',19),('Subtitle',11),('Heading 1',14),('Heading 2',11.5),('Caption',9)]:
    style=doc.styles[name]; style.font.name='Times New Roman'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
    style.font.bold=name.startswith('Heading'); style.font.underline=False
    fonts=style.element.get_or_add_rPr().rFonts
    if fonts is not None:
        for key in list(fonts.attrib):
            if 'theme' in key.lower(): del fonts.attrib[key]
    style.paragraph_format.space_after=Pt(5)
    style.paragraph_format.line_spacing=1.05
    if name.startswith('Heading'): style.paragraph_format.space_before=Pt(8)
normal=doc.styles['Normal']; normal.paragraph_format.widow_control=True
prompt=doc.styles.add_style('Author prompt',1)
prompt.base_style=normal; prompt.font.italic=True; prompt.font.size=Pt(9.5); prompt.font.color.rgb=RGBColor.from_string('666666')
doc.styles['Caption'].font.italic=True
doc.core_properties.title='Jet Bridge Air Speed Reconstruction Report Template'
doc.core_properties.subject='MMA3001 report structure with verified project results and author prompts'
doc.core_properties.author='[Student name]'
doc.core_properties.keywords='MMA3001 interpolation SVR CFD report template'
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('MMA3001  |  Page ')
for instruction in ['PAGE']:
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),instruction)
    footer._p.append(field)
for r in footer.runs: r.font.size=Pt(8)

def p(text,style=None): return doc.add_paragraph(text,style)
def h(text,level=1): return doc.add_heading(text,level)
def note(text): return p('[AUTHOR TASK  '+text+']','Author prompt')
def page(title):
    para=h(title)
    para.paragraph_format.page_break_before=True
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,v in enumerate(headers): t.rows[0].cells[i].text=v
    for row in rows:
        for cell,v in zip(t.add_row().cells,row): cell.text=str(v)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        if ri==0: trPr.append(OxmlElement('w:tblHeader'))
        no_split=OxmlElement('w:cantSplit'); trPr.append(no_split)
        for ci,cell in enumerate(row.cells):
            if widths: cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
            pr.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E6EDF2' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF')); pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'65'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=1.0
                if ci>0: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs: run.font.size=Pt(9); run.bold=ri==0
    p('')
    return t

figure=0
def fig(relative,caption,width=6.55):
    global figure
    figure+=1
    para=p(''); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.keep_with_next=True
    run=para.add_run(); picture=run.add_picture(str(ROOT/relative),width=Inches(width))
    picture._inline.docPr.set('descr',caption)
    c=p(f'Figure {figure}. {caption}','Caption'); c.paragraph_format.keep_together=True
    return figure


